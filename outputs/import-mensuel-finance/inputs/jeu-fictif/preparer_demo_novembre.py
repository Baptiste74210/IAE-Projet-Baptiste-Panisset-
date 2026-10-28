"""Prépare le mois de démonstration « Novembre 2026 » (Run du 2026-10-29) — données entièrement fictives.

À lancer depuis la racine du dépôt :
  python3 -I outputs/import-mensuel-finance/inputs/jeu-fictif/preparer_demo_novembre.py

Le mois de démo est le mois qui suit le jeu de référence d'octobre :
  - même database centrale fictive (400 contrats, inchangée) ;
  - même liste de contrats que le fichier finance d'octobre (mêmes pièges : 14 WBS mal saisis, faux ami,
    doublons, lignes sans WBS, multiservices sans code, nouveaux contrats PBH, N/A, types différents),
    avec de nouveaux montants ;
  - 3 nouveautés de novembre : 2 WBS nouvellement mal saisis (une inversion de deux caractères voisins,
    une lettre remplacée) et 1 nouveau contrat PBH inconnu de la database ;
  - mémoire de départ = décisions justes d'octobre (14 paires confirmées, faux ami rejeté).
Ce que la démo doit montrer : les 14 WBS confirmés en octobre passent à 100 % tout seuls, le faux ami
n'est plus proposé, et seules les 2 nouvelles erreurs de novembre sont demandées à Baptiste.
Graine fixe : mêmes fichiers à chaque lancement. Les fichiers d'entrée d'octobre ne sont jamais modifiés.
"""
import os, random, shutil
import openpyxl

JEU = os.path.join("outputs", "import-mensuel-finance", "inputs", "jeu-fictif")
OUT = os.path.join(JEU, "demo")
SRC = os.path.join(JEU, "Xbh Contracts Financial performance - Octobre 2026 (FICTIF).xlsx")
MEM = os.path.join(JEU, "memoire-E6 (FICTIF).csv")
DB = os.path.join(JEU, "Database Hcare XBH in service contracts (FICTIF).xlsx")
rng = random.Random(202611)


def main():
    os.makedirs(OUT, exist_ok=True)
    wb = openpyxl.load_workbook(SRC)
    ws = wb.active
    ws["A1"] = "Xbh Contracts Financial performance : Novembre 2026 — DONNÉES FICTIVES (démo)"
    db_keys = {str(r[0]).strip().lower() for r in openpyxl.load_workbook(DB).active.iter_rows(
        min_row=2, values_only=True) if r[0]}
    fin_keys = {}
    for r in range(20, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k:
            fin_keys.setdefault(str(k).strip().lower(), []).append(r)
    # Nouveaux montants (± 8 % autour d'octobre), N/A conservés tels quels.
    for r in range(20, ws.max_row + 1):
        for c in (6, 7, 8):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                ws.cell(r, c).value = round(v * rng.uniform(0.92, 1.08), 2)
        v = ws.cell(r, 9).value
        if isinstance(v, (int, float)):
            ws.cell(r, 9).value = max(0, int(round(v * rng.uniform(0.92, 1.08))))
    # Deux WBS XBH exacts en octobre, mal saisis en novembre.
    exacts = [r for r in range(20, ws.max_row + 1)
              if ws.cell(r, 3).value == "XBH" and isinstance(ws.cell(r, 1).value, str)
              and ws.cell(r, 1).value == ws.cell(r, 1).value.strip().upper()
              and ws.cell(r, 1).value.lower() in db_keys and len(fin_keys[ws.cell(r, 1).value.lower()]) == 1]
    r1, r2 = exacts[40], exacts[120]
    k1 = ws.cell(r1, 1).value
    k1_new = k1[:4] + k1[5] + k1[4] + k1[6:]          # inversion de deux caractères voisins
    k2 = ws.cell(r2, 1).value
    k2_new = k2[:6] + ("K" if k2[6] != "K" else "M") + k2[7:]   # une lettre remplacée
    ws.cell(r1, 1).value = k1_new
    ws.cell(r2, 1).value = k2_new
    # Nouveau contrat PBH inconnu, ajouté en fin de liste (le bloc de totaux est étendu).
    last = ws.max_row + 1
    ws.cell(last, 1).value = "FH.N7VEA"
    ws.cell(last, 2).value = "Client 401"
    ws.cell(last, 3).value = "PBH"
    ws.cell(last, 4).value = "Europe"
    ws.cell(last, 5).value = "France"
    ws.cell(last, 6).value = 2150400.00
    ws.cell(last, 7).value = 2098120.50
    ws.cell(last, 8).value = 1611800.25
    ws.cell(last, 9).value = 1840
    for r in range(5, 18):
        f = ws.cell(r, 6).value
        if isinstance(f, str) and f.startswith("="):
            ws.cell(r, 6).value = f.replace(str(last - 1), str(last))
    name = "Xbh Contracts Financial performance - Novembre 2026 (FICTIF - démo).xlsx"
    wb.save(os.path.join(OUT, name))
    shutil.copyfile(MEM, os.path.join(OUT, "memoire-apres-octobre (FICTIF).csv"))
    notes = [
        "# Corrigé de la démo Novembre 2026 — à ne pas donner au programme", "",
        f"Fichier : `demo/{name}` — lignes de contrats 20 à {last}. Mémoire de départ : "
        "`demo/memoire-apres-octobre (FICTIF).csv` (décisions justes d'octobre).", "",
        "## Nouveautés de novembre", "",
        f"| Ligne | Clé finance | Clé database | Piège | Bonne décision |", "|---|---|---|---|---|",
        f"| {r1} | `{k1_new}` | `{k1}` | deux caractères voisins inversés | **Oui** |",
        f"| {r2} | `{k2_new}` | `{k2}` | une lettre remplacée | **Oui** |",
        f"| {last} | `FH.N7VEA` | — | nouveau contrat PBH inconnu | (a), à traiter à la main |", "",
        "## Ce que la démo doit montrer", "",
        "- les 14 WBS confirmés en octobre sont rapprochés à 100 % grâce à la mémoire, sans question ;",
        "- le faux ami rejeté en octobre (`FH.F4ZSJ`) n'est plus proposé : il part en (a) ;",
        "- seules les 2 nouvelles erreurs de novembre sont à confirmer ;",
        "- les totaux restent `OK`.", "",
        "## Comptages attendus (vérifiés en lançant le programme une fois à la préparation)", "",
        "| 100 % | À confirmer (Élevée / Moyenne / Faible) | Non rapprochées | Total |", "|---|---|---|---|",
        f"| 382 | 2 (1 / 1 / 0) | 15 | {last - 19} |", "",
        "Après les 2 « Oui » : 384 / 0 / 15 ; totaux `OK` avant et après.",
    ]
    with open(os.path.join(OUT, "corrige-demo-novembre.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(notes) + "\n")
    print("\n".join(notes))


if __name__ == "__main__":
    main()
