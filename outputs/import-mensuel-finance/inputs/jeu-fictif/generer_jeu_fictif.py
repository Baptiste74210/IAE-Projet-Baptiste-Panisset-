"""Génère le jeu fictif de référence (400 contrats) et son corrigé — données entièrement inventées.

Aucun nom de client : la colonne client ne contient que des codes neutres (« Client 001 »…), et les
contrats multiservices sont identifiés par un code `MS-…` écrit dans la colonne WBS des deux fichiers.

À lancer depuis la racine du dépôt :
  python3 -I outputs/import-mensuel-finance/inputs/jeu-fictif/generer_jeu_fictif.py

Le corrigé est construit à partir des pièges placés par ce générateur, et non en exécutant le programme
testé : il reste donc une référence indépendante. Le tirage est fixé (graine 2026) : relancer le script
produit exactement les mêmes fichiers.
"""
import csv, json, os, random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

JEU = os.path.join("outputs", "import-mensuel-finance", "inputs", "jeu-fictif")
INPUTS = os.path.dirname(JEU)
MOIS = "Octobre 2026"
N_DB = 400
rng = random.Random(2026)

REGIONS = {"Europe": ["France", "Allemagne", "Italie", "Espagne", "Norvège", "Suisse", "Belgique", "Pologne"],
           "NORAM": ["États-Unis", "Canada"],
           "LATAM": ["Brésil", "Mexique", "Chili", "Colombie", "Argentine"],
           "APAC": ["Japon", "Australie", "Singapour", "Inde", "Corée du Sud", "Malaisie"],
           "AFME": ["Afrique du Sud", "Émirats arabes unis", "Nigeria", "Maroc", "Arabie saoudite", "Kenya"]}
MOTS1 = ["Alpine", "Atlas", "Aurora", "Blue", "Boreal", "Cardinal", "Cedar", "Coastal", "Condor", "Crystal",
         "Delta", "Eagle", "Falcon", "Golden", "Granite", "Harbor", "Horizon", "Ibex", "Jade", "Kestrel",
         "Lynx", "Meridian", "Nordic", "Nova", "Onyx", "Orion", "Pacific", "Polar", "Quartz", "Raven",
         "Sierra", "Solar", "Summit", "Titan", "Ultra", "Vector", "Vista", "Wind", "Zenith", "Zephyr"]
MOTS2 = ["Aviation", "Air", "Heli", "Rotor", "Wings", "Aero", "Flight", "Sky", "Lift", "Fleet"]
MOTS3 = ["Services", "Group", "Operations", "Rescue", "Medevac", "Offshore", "Patrol", "Charter", "Logistics", "Systems"]
ALPHA = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # sans O, I, 0, 1 : les confusions typiques sont réservées aux pièges


def lev(a, b):
    p = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        c = [i]
        for j, cb in enumerate(b, 1):
            c.append(min(p[j] + 1, c[j - 1] + 1, p[j - 1] + (ca != cb)))
        p = c
    return p[-1]


def nouvelle_cle(existantes, motif=None):
    while True:
        corps = motif() if motif else "".join(rng.choice(ALPHA) for _ in range(5))
        k = "FH." + corps
        if all(lev(k, e) >= 3 for e in existantes):
            return k


def clients_uniques(n):
    vus, out = set(), []
    while len(out) < n:
        nom = f"{rng.choice(MOTS1)} {rng.choice(MOTS2)} {rng.choice(MOTS3)}"
        if nom not in vus:
            vus.add(nom)
            out.append(nom)
    return out


def kpi(type_):
    to = round(rng.uniform(250_000, 7_500_000), 2)
    inv = round(to * rng.uniform(0.93, 1.0), 2)
    cost = round(to * rng.uniform(0.6, 0.86), 2)
    fh = 0 if type_ == "Multiservices" else rng.randint(300, 9_000)
    return [to, inv, cost, fh]


# ---------------------------------------------------------------- database centrale (400 contrats)
types = ["XBH"] * 222 + ["FBH"] * 88 + ["PBH"] * 80 + ["Multiservices"] * 10
rng.shuffle(types)
clients = [f"Client {i + 1:03d}" for i in range(N_DB)]   # codes neutres, aucun nom de client
db, cles = [], []
for i, t in enumerate(types):
    reg = rng.choice(list(REGIONS))
    c = {"id": i, "client": clients[i], "type": t, "region": reg, "pays": rng.choice(REGIONS[reg]), "wbs": ""}
    if t != "Multiservices":
        c["wbs"] = nouvelle_cle(cles)
        cles.append(c["wbs"])
    db.append(c)

par_type = {t: [c for c in db if c["type"] == t] for t in ("XBH", "FBH", "PBH", "Multiservices")}
pris = set()


def choisir(t, n):
    libres = [c for c in par_type[t] if c["id"] not in pris]
    sel = rng.sample(libres, n)
    pris.update(c["id"] for c in sel)
    return sel


def remplace_cle(c, nouvelle):
    cles.remove(c["wbs"])
    c["wbs"] = nouvelle
    cles.append(nouvelle)


# Pièges côté database / finance (types choisis pour que le corrigé soit certain, voir README du corrigé).
absents_finance = choisir("XBH", 8)                       # (b)
sans_wbs = choisir("XBH", 3)                              # finance sans WBS → (d) ; database → (b)
doublons = choisir("XBH", 2)                              # finance en double → (e)
na = choisir("XBH", 3)                                    # Cost = « N/A » → (f), rapproché
espaces = choisir("XBH", 6) + choisir("FBH", 4)           # espaces, majuscules, insécables → 100 %
type_diff = choisir("PBH", 2)                             # même WBS, type finance « XBH » → (h), rapproché
faux_amis_cibles = choisir("FBH", 1)                      # absent du fichier finance ; un nouveau contrat lui ressemble
ms = par_type["Multiservices"]
ms_sans_code = ms[:2]                                     # pas encore codés : (c) des deux côtés
ms_codes = ms[2:]
for i, c in enumerate(ms_codes, 1):                       # code MS-… écrit dans la colonne WBS des deux fichiers
    c["wbs"] = f"MS-{i:03d}"
pris.update(c["id"] for c in ms)

# WBS mal saisis côté finance : on fabrique une clé database qui contient le caractère à confondre.
divergents = []
def avec_motif(motif):
    return lambda: motif
specs = [("O/0", "XBH"), ("O/0", "XBH"), ("O/0", "FBH"), ("I/1", "XBH"), ("I/1", "XBH"), ("I/1", "FBH"),
         ("préfixe FH/FB", "XBH"), ("préfixe FH/FB", "FBH"), ("inversion", "XBH"), ("inversion", "FBH"),
         ("préfixe + inversion", "XBH"), ("préfixe + inversion", "FBH"), ("substitution", "XBH"), ("substitution", "FBH")]
for nature, t in specs:
    c = choisir(t, 1)[0]
    autres = [k for k in cles if k != c["wbs"]]
    while True:
        corps = list("".join(rng.choice(ALPHA) for _ in range(5)))
        if nature == "O/0":
            corps[rng.randrange(5)] = "0"
        elif nature == "I/1":
            corps[rng.randrange(5)] = "1"
        cible = "FH." + "".join(corps)
        p = rng.randrange(4)
        if nature == "O/0":
            p = corps.index("0"); f = list(corps); f[p] = "O"; saisie = "FH." + "".join(f)
        elif nature == "I/1":
            p = corps.index("1"); f = list(corps); f[p] = "I"; saisie = "FH." + "".join(f)
        elif nature == "préfixe FH/FB":
            saisie = "FB." + "".join(corps)
        elif nature in ("inversion", "préfixe + inversion"):
            if corps[p] == corps[p + 1]:
                continue
            f = list(corps); f[p], f[p + 1] = f[p + 1], f[p]
            saisie = ("FB." if nature.startswith("préfixe") else "FH.") + "".join(f)
        else:
            f = list(corps); f[p] = rng.choice([x for x in ALPHA if x != corps[p]]); saisie = "FH." + "".join(f)
        if all(lev(cible, k) >= 3 for k in autres) and all(lev(saisie, k) >= 3 for k in autres) and saisie != cible:
            break
    remplace_cle(c, cible)
    niveau = "Élevée" if nature in ("O/0", "I/1", "préfixe FH/FB", "inversion") else "Moyenne"
    divergents.append({"db": c, "saisie": saisie, "nature": nature, "niveau": niveau})

# Faux amis : nouveaux contrats FBH dont le WBS diffère d'un seul caractère (non typique) d'un contrat absent.
faux_amis = []
for cible in faux_amis_cibles:
    autres = [k for k in cles if k != cible["wbs"]]
    while True:
        p = rng.randrange(3, 8)
        s = list(cible["wbs"]); s[p] = rng.choice([x for x in ALPHA if x != s[p]]); saisie = "".join(s)
        if all(lev(saisie, k) >= 3 for k in autres):
            break
    reg = rng.choice(list(REGIONS))
    faux_amis.append({"cible": cible, "saisie": saisie, "client": f"Client {N_DB + 1:03d}",
                      "region": reg, "pays": rng.choice(REGIONS[reg])})
    cles.append(saisie)

# Nouveaux contrats PBH (aucun contrat PBH libre dans la database) → (a).
nouveaux = []
for n_ in range(4):
    reg = rng.choice(list(REGIONS))
    k = nouvelle_cle(cles); cles.append(k)
    nouveaux.append({"wbs": k, "client": f"Client {N_DB + 2 + n_:03d}",
                     "region": reg, "pays": rng.choice(REGIONS[reg])})

# ---------------------------------------------------------------- fichier finance
ids_absents = {c["id"] for c in absents_finance} | {c["id"] for c in faux_amis_cibles}
saisie_div = {d["db"]["id"]: d["saisie"] for d in divergents}
esp_formes = [lambda k: k + " ", lambda k: " " + k, lambda k: k + "  ", lambda k: " " + k,
              lambda k: k + " ", lambda k: k.lower(), lambda k: k[:3] + k[3:].lower(), lambda k: k + " ",
              lambda k: "​" + k, lambda k: " " + k + " "]
esp = {c["id"]: esp_formes[i] for i, c in enumerate(espaces)}
fin = []
for c in db:
    if c["id"] in ids_absents:
        continue
    k = kpi(c["type"])
    wbs = c["wbs"]
    if c["id"] in saisie_div:
        wbs = saisie_div[c["id"]]
    elif c["id"] in esp:
        wbs = esp[c["id"]](wbs)
    elif c["id"] in {x["id"] for x in sans_wbs}:
        wbs = ""
    if c["id"] in {x["id"] for x in na}:
        k[2] = "N/A"
    typ = "XBH" if c["id"] in {x["id"] for x in type_diff} else c["type"]
    fin.append({"src": c["id"], "wbs": wbs, "client": c["client"], "type": typ, "region": c["region"], "pays": c["pays"], "kpi": k})
    if c["id"] in {x["id"] for x in doublons}:
        fin.append({"src": c["id"], "wbs": c["wbs"], "client": c["client"], "type": c["type"], "region": c["region"],
                    "pays": c["pays"], "kpi": kpi(c["type"])})
for f in faux_amis:
    fin.append({"src": None, "wbs": f["saisie"], "client": f["client"], "type": "FBH", "region": f["region"], "pays": f["pays"], "kpi": kpi("FBH")})
for n in nouveaux:
    fin.append({"src": None, "wbs": n["wbs"], "client": n["client"], "type": "PBH", "region": n["region"], "pays": n["pays"], "kpi": kpi("PBH")})
rng.shuffle(fin)
for i, r in enumerate(fin):
    r["ligne"] = 20 + i

# ---------------------------------------------------------------- écriture des fichiers
F = "Arial"
hf = Font(name=F, bold=True, color="FFFFFF"); hp = PatternFill("solid", fgColor="1F2A5C")
th = Side(style="thin", color="C9CED8"); bx = Border(left=th, right=th, top=th, bottom=th); eur = '#,##0.00 "€"'
FH_ = ["WBS", "Client Name", "Contract Type", "Region", "Country", "Turnover (€)", "Invoicing (€)", "Cost (€)", "Flight Hours"]
DH_ = ["WBS", "Client", "Type de contrat", "Région", "Pays", "Turnover", "Invoicing", "Cost", "FH"]
first, last = 20, 19 + len(fin)

wb = Workbook(); ws = wb.active; ws.title = "XBH Contracts"
ws.merge_cells("A1:I1"); ws["A1"] = f"Xbh Contracts Financial performance : {MOIS} — DONNÉES FICTIVES"
ws["A1"].font = Font(name=F, bold=True, size=14); ws["A1"].alignment = Alignment(horizontal="center")
ws.merge_cells("A2:I2"); ws["A2"] = "Jeu de test fictif — aucune donnée Airbus réelle"; ws["A2"].font = Font(name=F, italic=True, color="888888")
ws.merge_cells("A4:E4"); ws["A4"] = "TOTAUX GLOBAUX — tous contrats"; ws["A4"].font = Font(name=F, bold=True)
for k, (lab, col) in enumerate([("Turnover (€)", "F"), ("Invoicing (€)", "G"), ("Cost (€)", "H"), ("Flight Hours", "I")]):
    r = 5 + k
    ws.merge_cells(f"A{r}:E{r}"); ws[f"A{r}"] = lab; ws[f"A{r}"].font = Font(name=F)
    ws[f"F{r}"] = f"=SUM({col}{first}:{col}{last})"; ws[f"F{r}"].font = Font(name=F, bold=True)
    ws[f"F{r}"].number_format = eur if col != "I" else "#,##0"
ws.merge_cells("A10:E10"); ws["A10"] = "Nombre de lignes de contrats"; ws["A10"].font = Font(name=F)
ws["F10"] = f"=COUNTA(C{first}:C{last})"; ws["F10"].font = Font(name=F, bold=True)
ws.merge_cells("A12:E12"); ws["A12"] = "TOTAUX PAR RÉGION — Turnover (€)"; ws["A12"].font = Font(name=F, bold=True)
for k, reg in enumerate(REGIONS):
    r = 13 + k
    ws.merge_cells(f"A{r}:E{r}"); ws[f"A{r}"] = reg; ws[f"A{r}"].font = Font(name=F)
    ws[f"F{r}"] = f'=SUMIFS(F{first}:F{last},D{first}:D{last},"{reg}")'; ws[f"F{r}"].number_format = eur
for c_, h in enumerate(FH_, 1):
    x = ws.cell(row=19, column=c_, value=h); x.font = hf; x.fill = hp; x.border = bx
for r in fin:
    vals = [r["wbs"] or None, r["client"], r["type"], r["region"], r["pays"]] + r["kpi"]
    for c_, v in enumerate(vals, 1):
        x = ws.cell(row=r["ligne"], column=c_, value=v); x.font = Font(name=F); x.border = bx
        if c_ in (6, 7, 8) and isinstance(v, float): x.number_format = eur
        if c_ == 9: x.number_format = "#,##0"
for col, w in zip("ABCDEFGHI", (14, 34, 15, 10, 20, 18, 18, 18, 13)): ws.column_dimensions[col].width = w
ws.freeze_panes = "A20"
wb.save(os.path.join(JEU, f"Xbh Contracts Financial performance - {MOIS} (FICTIF).xlsx"))

wb2 = Workbook(); w2 = wb2.active; w2.title = "Contracts"
for c_, h in enumerate(DH_, 1):
    x = w2.cell(row=1, column=c_, value=h); x.font = hf; x.fill = hp; x.border = bx
for i, c in enumerate(db):
    for c_, v in enumerate([c["wbs"] or None, c["client"], c["type"], c["region"], c["pays"]], 1):
        x = w2.cell(row=2 + i, column=c_, value=v); x.font = Font(name=F); x.border = bx
    for c_ in range(6, 10): w2.cell(row=2 + i, column=c_).border = bx
for col, w in zip("ABCDEFGHI", (14, 34, 15, 10, 20, 13, 13, 13, 10)): w2.column_dimensions[col].width = w
w2.freeze_panes = "A2"
wb2.save(os.path.join(JEU, "Database Hcare XBH in service contracts (FICTIF).xlsx"))

# Référentiel des contrats multiservices : aide-mémoire de Baptiste, jamais lu par le programme.
with open(os.path.join(JEU, "referentiel-multiservices (FICTIF).csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["Code MS", "Désignation interne du contrat", "Statut"])
    for c in ms_codes:
        w.writerow([c["wbs"], f"Contrat multiservices {c['wbs'][3:]} (fictif)", "actif"])
    w.writerow([f"MS-{len(ms_codes) + 1:03d}", f"Contrat multiservices {len(ms_codes) + 1:03d} (fictif)", "clos"])

with open(os.path.join(JEU, "table-correspondance-colonnes (FICTIF).csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["Colonne finance", "Colonne database", "Cellule du total"])
    for s, t, cell in [("Turnover (€)", "Turnover", "F5"), ("Invoicing (€)", "Invoicing", "F6"),
                       ("Cost (€)", "Cost", "F7"), ("Flight Hours", "FH", "F8")]:
        w.writerow([s, t, cell])

memoire_e6 = [(d["saisie"], d["db"]["wbs"], "confirmée") for d in divergents] + \
             [(f["saisie"], f["cible"]["wbs"], "rejetée") for f in faux_amis]
with open(os.path.join(JEU, "memoire-E6 (FICTIF).csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["clé finance", "clé database", "statut", "date"])
    for a, b, s in memoire_e6:
        w.writerow([a, b, s, "2026-10-27"])

# ---------------------------------------------------------------- corrigé (construit à partir des pièges)
def somme(lignes, i):
    return round(sum(r["kpi"][i] for r in lignes if isinstance(r["kpi"][i], (int, float))), 2)

div_src = {d["db"]["id"] for d in divergents}
dbl_src = {c["id"] for c in doublons}
sans_src = {c["id"] for c in sans_wbs}
ms_abs_src = {c["id"] for c in ms_sans_code}
fa_keys = {f["saisie"] for f in faux_amis}
new_keys = {n["wbs"] for n in nouveaux}
non_rapp_e5 = [r for r in fin if r["src"] in dbl_src or r["src"] in sans_src or r["src"] in ms_abs_src or r["wbs"] in new_keys]
attente_e5 = [r for r in fin if r["src"] in div_src or r["wbs"] in fa_keys]
non_rapp_e6 = non_rapp_e5 + [r for r in fin if r["wbs"] in fa_keys]
total = len(fin)
cent_e5 = total - len(attente_e5) - len(non_rapp_e5)
kp = ["Turnover (€)", "Invoicing (€)", "Cost (€)", "Flight Hours"]
bloc = {h: somme(fin, i) for i, h in enumerate(kp)}
attendu = {
    "mois": MOIS, "lignes_finance": total, "contrats_database": N_DB,
    "E5": {"100 %": cent_e5, "En attente – Élevée": sum(d["niveau"] == "Élevée" for d in divergents),
           "En attente – Moyenne": sum(d["niveau"] == "Moyenne" for d in divergents) + len(faux_amis),
           "En attente – Faible": 0, "Non rapprochées": len(non_rapp_e5), "Total lignes fichier finance": total},
    "E6": {"100 %": cent_e5 + len(divergents), "En attente – Élevée": 0, "En attente – Moyenne": 0,
           "En attente – Faible": 0, "Non rapprochées": len(non_rapp_e6), "Total lignes fichier finance": total},
    "bloc": bloc,
    "database_E5": {h: round(bloc[h] - somme(non_rapp_e5, i), 2) for i, h in enumerate(kp)},
    "database_E6": {h: round(bloc[h] - somme(non_rapp_e6, i), 2) for i, h in enumerate(kp)},
    "a_confirmer_E5": [{"cle_finance": d["saisie"], "cle_db": d["db"]["wbs"], "niveau": d["niveau"], "nature": d["nature"],
                        "bonne_decision": "Oui"} for d in divergents]
                      + [{"cle_finance": f["saisie"], "cle_db": f["cible"]["wbs"], "niveau": "Moyenne",
                          "nature": "faux ami (contrats différents)", "bonne_decision": "Non"} for f in faux_amis],
    "raisons_E5": {"(a) absent de la database": len(nouveaux), "(b) absent du fichier finance": len(absents_finance) + len(sans_wbs),
                   "(c) multiservices sans code": 2 * len(ms_sans_code), "(d) clé manquante": len(sans_wbs),
                   "(e) doublon": 2 * len(doublons) + len(doublons), "(f) valeur non numérique": len(na),
                   "(h) type de contrat différent": len(type_diff)},
    "multiservices": {"codés MS-… (100 %)": len(ms_codes), "sans code (c)": len(ms_sans_code)},
}
with open(os.path.join(JEU, "corrige-jeu-fictif.json"), "w", encoding="utf-8") as f:
    json.dump(attendu, f, ensure_ascii=False, indent=1)


def fe(v): return f"{v:,.2f}".replace(",", " ").replace(".", ",")
def fh(v): return f"{v:,.0f}".replace(",", " ")
def fmt(h, v): return fh(v) if h == "Flight Hours" else fe(v)

L = ["# Corrigé du jeu fictif de référence (400 contrats) — à ne pas donner au programme",
     "", "*Généré par `generer_jeu_fictif.py` à partir des pièges qu'il a placés (graine fixe 2026). Version machine : `corrige-jeu-fictif.json`. "
     "Ne jamais fournir ce fichier à une exécution : il sert uniquement à vérifier les résultats.*", "",
     f"Fichier finance : **{total} lignes** de contrats (lignes 20 à {last}). Database centrale : **{N_DB} contrats**. Mois : **{MOIS}**.", "",
     "## Pièges placés", "",
     "| Piège | Nombre | Résultat attendu (E5, mémoire vide) |", "|---|---|---|",
     f"| WBS identiques | {cent_e5 - len(espaces) - len(na) - len(type_diff) - len(ms_codes)} | Rapprochés à **100 %** |",
     f"| Espaces, espaces insécables, caractère invisible, minuscules | {len(espaces)} | Rapprochés à **100 %** |",
     f"| Multiservices codés `MS-…` dans la colonne WBS | {len(ms_codes)} | Rapprochés à **100 %** par leur code |",
     f"| Cost = « N/A » | {len(na)} | Rapprochés à **100 %**, signalés (f) |",
     f"| Même WBS, type différent (finance XBH, database PBH) | {len(type_diff)} | Rapprochés à **100 %**, signalés (h) |",
     f"| WBS mal saisis (O/0, I/1, préfixe FH/FB, inversion) | {sum(d['niveau'] == 'Élevée' for d in divergents)} | À confirmer, confiance **Élevée** ; bonne décision **Oui** |",
     f"| WBS mal saisis (deux erreurs typiques, ou une autre lettre) | {sum(d['niveau'] == 'Moyenne' for d in divergents)} | À confirmer, confiance **Moyenne** ; bonne décision **Oui** |",
     f"| Faux ami : nouveau contrat à un caractère d'un contrat absent | {len(faux_amis)} | À confirmer, confiance **Moyenne** ; bonne décision **Non** |",
     f"| Doublons (même WBS sur deux lignes) | {len(doublons)} WBS ({2 * len(doublons)} lignes) | Non rapprochés, (e) ; jamais additionnés |",
     f"| Lignes sans WBS (type XBH) | {len(sans_wbs)} | Non rapprochées, (d) ; jamais rapprochées par le nom |",
     f"| Multiservices sans code (WBS vide) | {len(ms_sans_code)} | Non rapprochés, (c), des deux côtés |",
     f"| Nouveaux contrats PBH | {len(nouveaux)} | Non rapprochés, (a) : aucun contrat PBH libre |",
     f"| Contrats de la database absents du fichier finance | {len(absents_finance)} | (b) |", "",
     "## Comptages attendus", "",
     "| Scénario | 100 % | En attente (Élevée / Moyenne / Faible) | Non rapprochées | Total |", "|---|---|---|---|---|",
     f"| **E5** — mémoire vide | {attendu['E5']['100 %']} | {attendu['E5']['En attente – Élevée']} / {attendu['E5']['En attente – Moyenne']} / 0 | {attendu['E5']['Non rapprochées']} | {total} |",
     f"| **E6** — mémoire `memoire-E6 (FICTIF).csv` | {attendu['E6']['100 %']} | 0 / 0 / 0 | {attendu['E6']['Non rapprochées']} | {total} |",
     f"| **Chaîne** — E5, décisions de la colonne « bonne décision », relance | {attendu['E6']['100 %']} après décisions | 0 / 0 / 0 | {attendu['E6']['Non rapprochées']} | {total} |", "",
     "Le scénario **E7 (multiservices)** se vérifie dans l'exécution de E5 : "
     f"{len(ms_codes)} multiservices rapprochés à 100 % par leur code `MS-…`, {len(ms_sans_code)} sans code en (c) des deux côtés ; la colonne client n'est jamais lue.", "",
     "## Totaux attendus", "",
     "| KPI | Bloc de totaux | Database E5 | Database E6 et chaîne |", "|---|---|---|---|"]
for h in kp:
    L.append(f"| {h} | {fmt(h, bloc[h])} | {fmt(h, attendu['database_E5'][h])} | {fmt(h, attendu['database_E6'][h])} |")
L += ["", "Dans tous les cas, l'écart restant est nul (statut `OK`) : database + lignes non rapprochées = bloc de totaux.", "",
      "## Correspondances à confirmer (E5)", "", "| Clé finance | Clé database | Nature | Confiance attendue | Bonne décision |", "|---|---|---|---|---|"]
for a in attendu["a_confirmer_E5"]:
    L.append(f"| `{a['cle_finance']}` | `{a['cle_db']}` | {a['nature']} | {a['niveau']} | **{a['bonne_decision']}** |")
L += ["", "## Pourquoi ce corrigé est certain", "",
      "- Toutes les clés de la database diffèrent d'au moins 3 caractères entre elles : un WBS mal saisi n'a qu'un seul candidat proche.",
      "- Les nouveaux contrats sont de type PBH et aucun contrat PBH de la database ne reste libre : ils ne peuvent être rapprochés de rien.",
      "- Le faux ami est de type FBH, et aucun autre contrat FBH ne reste libre une fois sa paire rejetée (scénario E6).",
      "- Les contrats absents du fichier finance et les lignes sans WBS sont de type XBH : les seules lignes XBH en recherche automatique sont les WBS mal saisis, qui trouvent leur contrat avant."]
open(os.path.join(JEU, "corrige-jeu-fictif.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")

# ---------------------------------------------------------------- fichier d'entrée lisible (E5)
md = [f"# E5 — Jeu fictif de référence : {total} lignes de contrats, {N_DB} contrats en database (proposed)", "",
      "**Données entièrement fictives** : WBS et montants inventés, aucune donnée Airbus ; **aucun nom de client** (colonne client remplie de codes neutres « Client 001 »…, jamais lue par le programme). "
      "Fichiers à utiliser (dans `inputs/jeu-fictif/`) :", "",
      f"- `Xbh Contracts Financial performance - {MOIS} (FICTIF).xlsx` — fichier finance : lignes 1 à 18 = bloc de totaux, ligne 19 = en-têtes, contrats à partir de la ligne 20 ;",
      "- `Database Hcare XBH in service contracts (FICTIF).xlsx` — database centrale, KPI vides ;",
      "- `referentiel-multiservices (FICTIF).csv` — référentiel des codes `MS-…` (aide-mémoire, jamais lu par le programme) ; les codes sont écrits dans la colonne WBS des deux fichiers ;",
      "- `table-correspondance-colonnes (FICTIF).csv` — colonnes de KPI et cellules des totaux ;",
      "- `memoire-E6 (FICTIF).csv` — mémoire des correspondances du scénario E6 (mois suivant).", "",
      f"Mois demandé au lancement : **{MOIS}**. Le contenu complet figure dans les fichiers Excel ; ce fichier n'en donne qu'un aperçu.", "",
      "## Aperçu des 15 premières lignes du fichier finance", "",
      "| Ligne | " + " | ".join(FH_) + " |", "|" + "---|" * (len(FH_) + 1)]
for r in fin[:15]:
    md.append(f"| {r['ligne']} | " + (f"`{r['wbs']}`" if r["wbs"] else "*(vide)*") + f" | {r['client']} | {r['type']} | {r['region']} | {r['pays']} | "
              + " | ".join(str(v) if not isinstance(v, float) else fe(v) for v in r["kpi"][:3]) + f" | {r['kpi'][3]} |")
open(os.path.join(INPUTS, "E5-jeu-fictif-reference.md"), "w", encoding="utf-8").write("\n".join(md) + "\n")

print(json.dumps({"lignes_finance": total, "E5": attendu["E5"], "E6": attendu["E6"]}, ensure_ascii=False))
