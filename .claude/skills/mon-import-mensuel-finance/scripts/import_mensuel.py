"""Simulation de l'import mensuel finance (données fictives uniquement).

À lancer depuis la racine du dépôt :
  Simulation : python3 -I .claude/skills/mon-import-mensuel-finance/scripts/import_mensuel.py simulation
               [--mois "Octobre 2026"] [--finance F.xlsx] [--database D.xlsx] [--pseudo P.csv] [--memoire M.csv]
  Décisions  : python3 -I .claude/skills/mon-import-mensuel-finance/scripts/import_mensuel.py decisions
               --fichier "outputs/import-mensuel-finance/runs/Database XBH Octobre 2026.xlsx"
               [--decision "FB.TEST060=Oui" --decision "FH.TEST011=Non"] [--memoire M.csv]
               (étape 8 : sans --decision, lit la colonne « Décision » remplie dans le fichier)
  Tests      : python3 -I .claude/skills/mon-import-mensuel-finance/scripts/import_mensuel.py tests [--garder]
               (E5, E6, E7, puis E5 → décisions → relance ; fichiers écrits dans un dossier temporaire,
                ou dans runs/ avec --garder)
Écrit la database du mois dans outputs/import-mensuel-finance/runs/ (nouvelle version si elle existe déjà)
et affiche un résumé JSON. Ne lit jamais le corrigé : la comparaison se fait après, à part.
Besoin : openpyxl, et LibreOffice (soffice) pour calculer les formules du bloc de totaux.
"""
import argparse, csv, json, os, re, shutil, subprocess, sys, tempfile, unicodedata
from datetime import date
import openpyxl

WORKFLOW = os.path.join("outputs", "import-mensuel-finance")
JEU = os.path.join(WORKFLOW, "inputs", "jeu-fictif")
RUNS = os.path.join(WORKFLOW, "runs")
DEFAULTS = {
    "finance": os.path.join(JEU, "Xbh Contracts Financial performance - Octobre 2026 (FICTIF).xlsx"),
    "database": os.path.join(JEU, "Database Hcare XBH in service contracts (FICTIF).xlsx"),
    "pseudo": os.path.join(JEU, "table-pseudonymisation-multiservices (FICTIF).csv"),
    "memoire": os.path.join(RUNS, "memoire-correspondances.csv"),
}
# Mémoire du scénario E6, telle que décrite dans requirements.md.
MEMOIRE_E6 = [("FB.TEST060", "FH.TEST006", "confirmée"), ("FH.TEST011", "FH.TEST010", "rejetée")]

MAP = {"Turnover (€)": "Turnover", "Invoicing (€)": "Invoicing", "Cost (€)": "Cost", "Flight Hours": "FH"}
KPI_SRC = list(MAP)
BLOCK_ROWS = {"Turnover (€)": 5, "Invoicing (€)": 6, "Cost (€)": 7, "Flight Hours": 8}
INVISIBLE = " ​‌‍﻿⁠ "
REASONS = {
    "a": "(a) absent de la database, sans aucun candidat",
    "b": "(b) absent du fichier finance",
    "c": "(c) multiservices sans code",
    "d": "(d) clé manquante",
    "e": "(e) doublon",
    "f": "(f) valeur non numérique",
    "g": "(g) correspondance rejetée",
    "h": "(h) type de contrat différent",
}


CONF_HEADERS = ["Clé finance", "Clé database", "Niveau de confiance", "Écart", "Décision (Oui / Non)",
                "Statut", "Ligne finance d'origine", "Ligne dans l'onglet Database"]
MEM_HEADERS = ["clé finance", "clé database", "statut", "date"]


class Stop(Exception):
    pass


def norm(k):
    if k is None:
        return ""
    s = str(k)
    s = s.strip(" \t\r\n" + INVISIBLE)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Cf")
    return s.upper()


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def sub_cost(a, b, i, j, s, t):
    if a == b:
        return 0
    if {a, b} in ({"O", "0"}, {"I", "1"}):
        return 0.5
    # lettre du préfixe FH/FB
    if i == 1 and j == 1 and s[0] == t[0] == "F" and {a, b} == {"H", "B"}:
        return 0.5
    return 1


def ecart(s, t):
    """Distance d'édition pondérée (Damerau restreinte) : 0,5 par erreur typique, 1 sinon."""
    n, m = len(s), len(t)
    d = [[0.0] * (m + 1) for _ in range(n + 1)]
    for i in range(n + 1):
        d[i][0] = i
    for j in range(m + 1):
        d[0][j] = j
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            d[i][j] = min(d[i - 1][j] + 1, d[i][j - 1] + 1,
                          d[i - 1][j - 1] + sub_cost(s[i - 1], t[j - 1], i - 1, j - 1, s, t))
            if i > 1 and j > 1 and s[i - 1] == t[j - 2] and s[i - 2] == t[j - 1] and s[i - 1] != s[i - 2]:
                d[i][j] = min(d[i][j], d[i - 2][j - 2] + 0.5)
    return d[n][m]


def load_pseudo(path):
    if not path or not os.path.exists(path):
        return None
    with open(path, newline="", encoding="utf-8-sig") as f:
        return {r["Client"].strip(): r["Code MS"].strip() for r in csv.DictReader(f)}


def load_memory(path):
    mem = {"confirmée": set(), "rejetée": set()}
    if path and path != "-" and os.path.exists(path):
        with open(path, newline="", encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                mem[r["statut"].strip()].add((norm(r["clé finance"]), norm(r["clé database"])))
    return mem


def anonymise(key, ctype, client, pseudo):
    """Renvoie (clé, raison éventuelle). Le nom du client ne sert qu'à la table."""
    k = norm(key)
    if k:
        return k, None
    if str(ctype or "").strip().lower() == "multiservices":
        if pseudo is not None and str(client or "").strip() in pseudo:
            return norm(pseudo[str(client).strip()]), None
        return "", "c"
    return "", "d"


def recalc_block(path):
    """Calcule les formules du bloc de totaux sur une copie (le fichier d'origine n'est pas touché)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise Stop("LibreOffice (soffice) est introuvable : impossible de calculer le bloc de totaux.")
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "finance.xlsx")
        shutil.copyfile(path, src)
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", os.path.join(tmp, "out"), src],
                       check=True, capture_output=True, timeout=180)
        ws = openpyxl.load_workbook(os.path.join(tmp, "out", "finance.xlsx"), data_only=True).active
        return {h: ws.cell(BLOCK_ROWS[h], 6).value for h in KPI_SRC}


def read_finance(path, mois):
    if mois.lower() not in os.path.basename(path).lower():
        raise Stop(f"Le nom du fichier finance ne contient pas le mois demandé « {mois} ».")
    wb = openpyxl.load_workbook(path)
    if len(wb.sheetnames) != 1:
        raise Stop(f"Le fichier finance contient {len(wb.sheetnames)} onglets au lieu d'un seul.")
    ws = wb.active
    headers = [c.value for c in ws[19]]
    col = {}
    for h in ["WBS", "Client Name", "Contract Type"] + KPI_SRC:
        n = headers.count(h)
        if n != 1:
            raise Stop(f"En-tête « {h} » trouvé {n} fois en ligne 19 (attendu : 1).")
        col[h] = headers.index(h)
    rows = []
    for r in range(20, ws.max_row + 1):
        vals = [c.value for c in ws[r]]
        if any(isinstance(v, str) and re.match(r"^\s*(sous-)?total", v, re.I) for v in vals):
            raise Stop(f"Ligne « Total » trouvée en ligne {r}.")
        if vals[col["WBS"]] in (None, "") and vals[col["Contract Type"]] in (None, ""):
            continue
        rows.append({"row": r, "wbs": vals[col["WBS"]], "client": vals[col["Client Name"]],
                     "type": vals[col["Contract Type"]], "kpi": {h: vals[col[h]] for h in KPI_SRC}})
    if not rows:
        raise Stop("Aucun contrat à partir de la ligne 20.")
    last = max(r["row"] for r in rows)
    rows = [r for r in rows if r["row"] <= last]
    block = recalc_block(path)
    for h, v in block.items():
        if not is_num(v):
            raise Stop(f"Total du bloc introuvable ou non numérique pour « {h} » (ligne {BLOCK_ROWS[h]}).")
    return rows, block


def read_db(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    for tgt in MAP.values():
        if tgt not in headers:
            raise Stop(f"Colonne « {tgt} » absente de la database centrale.")
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [c.value for c in ws[r]]
        if all(v in (None, "") for v in vals):
            continue
        rec = dict(zip(headers, vals))
        for tgt in MAP.values():
            if rec.get(tgt) not in (None, ""):
                raise Stop(f"KPI déjà rempli dans la database centrale : ligne {r}, colonne {tgt}.")
        rows.append({"row": r, "values": vals, "wbs": rec.get("WBS"), "client": rec.get("Client"),
                     "type": rec.get("Type de contrat")})
    return headers, rows


def run(fin_path, db_path, pseudo_path, mem_path, mois, runs_dir, label):
    for p in (fin_path, db_path):
        if not os.path.exists(p):
            raise Stop(f"Fichier d'entrée introuvable : {p}")
    pseudo = load_pseudo(pseudo_path)
    mem = load_memory(mem_path)
    fin, block = read_finance(fin_path, mois)
    headers, db = read_db(db_path)
    notes = [] if pseudo is not None else ["Table de pseudonymisation absente."]

    # Anonymisation : on ne garde que ligne, clé, type et KPI.
    anon_fin, anon_db = [], []
    for r in fin:
        k, why = anonymise(r["wbs"], r["type"], r["client"], pseudo)
        anon_fin.append({"row": r["row"], "key": k, "type": r["type"], "kpi": r["kpi"], "why": why})
    for r in db:
        k, why = anonymise(r["wbs"], r["type"], r["client"], pseudo)
        anon_db.append({"row": r["row"], "key": k, "type": r["type"], "why": why})

    fin_count, db_count = {}, {}
    for r in anon_fin:
        if r["key"]:
            fin_count[r["key"]] = fin_count.get(r["key"], 0) + 1
    for r in anon_db:
        if r["key"]:
            db_count[r["key"]] = db_count.get(r["key"], 0) + 1
    for r in anon_fin:
        if r["key"] and fin_count[r["key"]] > 1:
            r["why"] = "e"

    matches = {}  # ligne finance -> dict
    used_db = set()
    db_by_key = {r["key"]: r for r in anon_db if r["key"] and db_count[r["key"]] == 1}

    # Passe 1 : clé identique, unique des deux côtés ; ou paire confirmée en mémoire.
    for r in anon_fin:
        if r["why"]:
            continue
        target = None
        d = db_by_key.get(r["key"])
        if d and d["row"] not in used_db:
            target = d
        else:
            for (fk, dk) in mem["confirmée"]:
                if fk == r["key"] and dk in db_by_key and db_by_key[dk]["row"] not in used_db:
                    target = db_by_key[dk]
        if target:
            used_db.add(target["row"])
            matches[r["row"]] = {"db": target, "level": "100 %", "ecart": 0,
                                 "type_diff": norm(target["type"]) != norm(r["type"])}

    # Passe 2 : plus proche de même type, paires rejetées exclues, attribution par écart croissant.
    pending = [r for r in anon_fin if not r["why"] and r["row"] not in matches]
    pairs = []
    for r in pending:
        for d in anon_db:
            if (d["key"] and d["row"] not in used_db and norm(d["type"]) == norm(r["type"])
                    and (r["key"], d["key"]) not in mem["rejetée"]):
                pairs.append((ecart(r["key"], d["key"]), r["row"], d["row"], r, d))
    pairs.sort(key=lambda p: (p[0], p[1], p[2]))
    for e, fr, dr, r, d in pairs:
        if fr in matches or dr in used_db:
            continue
        free = [p[0] for p in pairs if p[1] == fr and p[2] not in used_db]
        best = min(free)
        at_best = free.count(best)
        if e == 0.5 and at_best == 1:
            level = "Élevée"
        elif e <= 2 and e == best and at_best == 1:
            level = "Moyenne"
        else:
            level = "Faible"
        used_db.add(dr)
        matches[fr] = {"db": d, "level": level, "ecart": e, "type_diff": False}
    for r in pending:
        if r["row"] not in matches:
            r["why"] = "a"

    # Lignes à traiter à la main.
    manual = []
    for r in anon_fin:
        m = matches.get(r["row"])
        if r["why"]:
            manual.append((r["row"], r["key"], r["type"], REASONS[r["why"]], r["kpi"]))
        elif m:
            bad = [h for h in KPI_SRC if not is_num(r["kpi"][h])]
            if bad:
                manual.append((r["row"], r["key"], r["type"],
                               REASONS["f"] + " : " + ", ".join(bad) + " (ligne rapprochée)", r["kpi"]))
            if m["type_diff"]:
                manual.append((r["row"], r["key"], r["type"],
                               REASONS["h"] + f" : database = {m['db']['type']} (ligne rapprochée)", r["kpi"]))
    dup_keys = {k for k, n in fin_count.items() if n > 1}
    for d in anon_db:
        if d["row"] in used_db:
            continue
        if d["key"] in dup_keys:
            why = REASONS["e"] + " côté finance (contrat de la database non rempli)"
        elif d["why"] == "c":
            why = REASONS["c"] + " (côté database)"
        else:
            why = REASONS["b"]
        manual.append((f"database l.{d['row']}", d["key"], d["type"], why, None))

    # Contrôle des totaux.
    fin_by_row = {r["row"]: r for r in anon_fin}
    control = []
    for h in KPI_SRC:
        lines = sum(r["kpi"][h] for r in anon_fin if is_num(r["kpi"][h]))
        dbt = sum(fin_by_row[fr]["kpi"][h] for fr in matches if is_num(fin_by_row[fr]["kpi"][h]))
        unm = sum(r["kpi"][h] for r in anon_fin if r["row"] not in matches and is_num(r["kpi"][h]))
        pend = sum(fin_by_row[fr]["kpi"][h] for fr, m in matches.items()
                   if m["level"] != "100 %" and is_num(fin_by_row[fr]["kpi"][h]))
        rest = round(block[h] - (dbt + unm), 2)
        ok = abs(block[h] - lines) <= 0.01 and abs(rest) <= 0.01
        control.append([h, block[h], lines, dbt, unm, pend, rest, "OK" if ok else "À vérifier"])
    counts = {"100 %": sum(1 for m in matches.values() if m["level"] == "100 %")}
    for lv in ("Élevée", "Moyenne", "Faible"):
        counts["En attente – " + lv] = sum(1 for m in matches.values() if m["level"] == lv)
    counts["Non rapprochées"] = len(anon_fin) - len(matches)
    counts["Total lignes fichier finance"] = len(anon_fin)
    assert sum(v for k, v in counts.items() if k != "Total lignes fichier finance") == len(anon_fin)

    # Écriture du fichier (jamais d'écrasement).
    os.makedirs(runs_dir, exist_ok=True)
    base = f"Database XBH {mois}"
    out, v = os.path.join(runs_dir, base + ".xlsx"), 1
    while os.path.exists(out):
        v += 1
        out = os.path.join(runs_dir, f"{base} v{v}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Database"
    ws.append(headers)
    by_db = {m["db"]["row"]: fin_by_row[fr] for fr, m in matches.items()}
    out_row = {}
    for d in db:  # informations des contrats remises depuis la database centrale
        out_row[d["row"]] = ws.max_row + 1
        vals = list(d["values"])
        if d["row"] in by_db:
            for src, tgt in MAP.items():
                vals[headers.index(tgt)] = by_db[d["row"]]["kpi"][src]
        ws.append(vals)
    ws2 = wb.create_sheet("À traiter à la main")
    ws2.append(["N° de ligne d'origine", "Clé", "Type de contrat", "Raison"]
               + [f"{h} (finance)" for h in KPI_SRC] + ["Valeur saisie à la main", "Date de saisie"])
    for row, key, typ, why, kpi in manual:
        ws2.append([row, key or "donnée manquante", typ, why]
                   + [kpi[h] if kpi else None for h in KPI_SRC] + [None, None])
    ws3 = wb.create_sheet("Correspondances à confirmer")
    ws3.append(CONF_HEADERS)
    to_confirm = []
    for fr, m in sorted(matches.items()):
        if m["level"] != "100 %":
            to_confirm.append((fin_by_row[fr]["key"], m["db"]["key"], m["level"], m["ecart"]))
            ws3.append([*to_confirm[-1], None, "En attente de confirmation", fr, out_row[m["db"]["row"]]])
    ws4 = wb.create_sheet("Contrôle")
    ws4.append(["KPI", "Total du bloc", "Total des lignes", "Total database", "Part non rapprochée",
                "Part en attente de confirmation", "Écart restant", "Statut"])
    for c in control:
        ws4.append(c)
    ws4.append([])
    for k, n in counts.items():
        ws4.append([k, n])
    ws4.append([])
    ws4.append(["Date d'exécution", date.today().isoformat()])
    ws4.append(["Fichier finance", os.path.basename(fin_path)])
    ws4.append(["Mois", mois])
    ws4.append(["Version", v])
    ws4.append(["Scénario", label])
    for n in notes:
        ws4.append(["Note", n])
    ws4.append([])
    ws4.append(["Validé", "indisponible : décisions à prendre" if to_confirm else "possible (case à cocher par Baptiste)"])
    wb.save(out)

    return {"scenario": label, "fichier": out, "counts": counts,
            "matches": [{"ligne": fr, "cle_finance": fin_by_row[fr]["key"], "cle_db": m["db"]["key"],
                         "niveau": m["level"], "ecart": m["ecart"]} for fr, m in sorted(matches.items())],
            "a_traiter": [{"ligne": str(a), "cle": b, "type": c, "raison": d} for a, b, c, d, _ in manual],
            "a_confirmer": to_confirm, "controle": control}


def apply_decisions(xlsx_path, mem_path, forced=None, today=None):
    """Étape 8 : applique les décisions Oui / Non de Baptiste, met à jour la mémoire et relance le contrôle.

    forced : dict {clé finance normalisée: "Oui"/"Non"} qui simule la saisie de Baptiste dans la colonne
    « Décision » ; sans forced, la colonne remplie dans le fichier fait foi. Seul le fichier du mois traité
    est mis à jour (c'est l'action prévue par l'étape 8) ; les fichiers d'entrée ne sont jamais touchés.
    """
    today = today or date.today().isoformat()
    if not os.path.exists(xlsx_path):
        raise Stop(f"Fichier du mois introuvable : {xlsx_path}")
    wb = openpyxl.load_workbook(xlsx_path)
    for name in ("Database", "À traiter à la main", "Correspondances à confirmer", "Contrôle"):
        if name not in wb.sheetnames:
            raise Stop(f"Onglet « {name} » absent : ce fichier n'a pas été produit par l'import.")
    wdb, wman, wconf, wctl = (wb[n] for n in ("Database", "À traiter à la main",
                                               "Correspondances à confirmer", "Contrôle"))
    if [c.value for c in wconf[1]][:len(CONF_HEADERS)] != CONF_HEADERS:
        raise Stop("Onglet « Correspondances à confirmer » d'un ancien format : relancer l'import d'abord.")
    headers = [c.value for c in wdb[1]]
    kcol = {src: headers.index(tgt) + 1 for src, tgt in MAP.items()}
    ci = {h: i + 1 for i, h in enumerate(CONF_HEADERS)}

    done = {"Oui": [], "Non": [], "vide": []}
    delta = {h: {"db": 0.0, "unm": 0.0, "pend": 0.0} for h in KPI_SRC}
    moved = {"100 %": 0, "Non rapprochées": 0}
    level_minus = {}
    new_mem = []
    for r in range(2, wconf.max_row + 1):
        fk = wconf.cell(r, ci["Clé finance"]).value
        if fk in (None, ""):
            continue
        if not str(wconf.cell(r, ci["Statut"]).value or "").startswith("En attente"):
            continue  # décision déjà prise lors d'un passage précédent
        if forced is not None and norm(fk) in forced:
            wconf.cell(r, ci["Décision (Oui / Non)"]).value = forced[norm(fk)]
        dec = str(wconf.cell(r, ci["Décision (Oui / Non)"]).value or "").strip().lower()
        dk = wconf.cell(r, ci["Clé database"]).value
        lvl = wconf.cell(r, ci["Niveau de confiance"]).value
        frow = wconf.cell(r, ci["Ligne finance d'origine"]).value
        drow = wconf.cell(r, ci["Ligne dans l'onglet Database"]).value
        vals = {h: wdb.cell(drow, kcol[h]).value for h in KPI_SRC}
        if dec not in ("oui", "non"):
            done["vide"].append(fk)
            continue
        for h in KPI_SRC:
            if is_num(vals[h]):
                delta[h]["pend"] -= vals[h]
        level_minus[lvl] = level_minus.get(lvl, 0) + 1
        if dec == "oui":
            wconf.cell(r, ci["Statut"]).value = f"Confirmée le {today} (certaine)"
            new_mem.append([norm(fk), norm(dk), "confirmée", today])
            moved["100 %"] += 1
            done["Oui"].append(fk)
        else:
            for h in KPI_SRC:
                if is_num(vals[h]):
                    delta[h]["db"] -= vals[h]
                    delta[h]["unm"] += vals[h]
                wdb.cell(drow, kcol[h]).value = None
            wconf.cell(r, ci["Statut"]).value = f"Rejetée le {today}"
            wman.append([frow, fk, None, REASONS["g"] + f" (proposée avec {dk})"]
                        + [vals[h] for h in KPI_SRC] + [None, None])
            wman.append([f"database l.{drow}", dk, None, REASONS["b"] + f" (après rejet de {fk})"]
                        + [None] * len(KPI_SRC) + [None, None])
            new_mem.append([norm(fk), norm(dk), "rejetée", today])
            moved["Non rapprochées"] += 1
            done["Non"].append(fk)

    # Mémoire des correspondances : ajout sans doublon, fichier créé s'il est absent.
    if new_mem and mem_path and mem_path != "-":
        existing = set()
        if os.path.exists(mem_path):
            with open(mem_path, newline="", encoding="utf-8-sig") as f:
                existing = {(norm(x["clé finance"]), norm(x["clé database"]), x["statut"].strip())
                            for x in csv.DictReader(f)}
        os.makedirs(os.path.dirname(mem_path) or ".", exist_ok=True)
        new_file = not os.path.exists(mem_path)
        with open(mem_path, "a", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            if new_file:
                w.writerow(MEM_HEADERS)
            for row in new_mem:
                if tuple(row[:3]) not in existing:
                    w.writerow(row)

    # Relance du contrôle (étape 5) : totaux et comptages mis à jour.
    statut = {}
    for r in range(2, wctl.max_row + 1):
        label = wctl.cell(r, 1).value
        if label in KPI_SRC:
            block, lines = wctl.cell(r, 2).value, wctl.cell(r, 3).value
            dbt = round(wctl.cell(r, 4).value + delta[label]["db"], 2)
            unm = round(wctl.cell(r, 5).value + delta[label]["unm"], 2)
            pend = round(wctl.cell(r, 6).value + delta[label]["pend"], 2)
            real_db = sum(v for v in (wdb.cell(i, kcol[label]).value for i in range(2, wdb.max_row + 1)) if is_num(v))
            if abs(real_db - dbt) > 0.01:
                raise Stop(f"Incohérence interne sur {label} : total database {real_db} ≠ {dbt}.")
            rest = round(block - (dbt + unm), 2)
            ok = abs(block - lines) <= 0.01 and abs(rest) <= 0.01
            for col, v in ((4, dbt), (5, unm), (6, pend), (7, rest), (8, "OK" if ok else "À vérifier")):
                wctl.cell(r, col).value = v
            statut[label] = "OK" if ok else "À vérifier"
        elif label == "100 %":
            wctl.cell(r, 2).value += moved["100 %"]
        elif isinstance(label, str) and label.startswith("En attente – "):
            wctl.cell(r, 2).value -= level_minus.get(label.split("– ", 1)[1], 0)
        elif label == "Non rapprochées":
            wctl.cell(r, 2).value += moved["Non rapprochées"]
        elif label == "Validé":
            wctl.cell(r, 2).value = ("indisponible : décisions à prendre" if done["vide"]
                                     else "possible (case à cocher par Baptiste)")
    wctl.append([])
    wctl.append(["Prise en compte des décisions", today])
    wctl.append(["Décisions « Oui »", len(done["Oui"])])
    wctl.append(["Décisions « Non »", len(done["Non"])])
    wctl.append(["Décisions manquantes", len(done["vide"])])
    wb.save(xlsx_path)
    counts = {}
    for r in range(2, wctl.max_row + 1):
        lab = wctl.cell(r, 1).value
        if lab in ("100 %", "Non rapprochées", "Total lignes fichier finance") or (
                isinstance(lab, str) and lab.startswith("En attente – ")):
            counts[lab] = wctl.cell(r, 2).value
    return {"fichier": xlsx_path, "oui": done["Oui"], "non": done["Non"], "manquantes": done["vide"],
            "memoire": mem_path, "counts": counts, "statut_controle": statut,
            "validation": "indisponible" if done["vide"] else "possible"}


def run_tests(mois, runs_dir=None):
    """E5 (mémoire vide), E6 (mémoire de requirements.md), E7 (multiservices) si ses fichiers existent."""
    results = []
    with tempfile.TemporaryDirectory() as tmp:
        mem_e6 = os.path.join(tmp, "memoire-E6.csv")
        with open(mem_e6, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["clé finance", "clé database", "statut", "date"])
            for fk, dk, st in MEMOIRE_E6:
                w.writerow([fk, dk, st, date.today().isoformat()])
        scenarios = [("E5 (mémoire vide)", DEFAULTS["finance"], DEFAULTS["database"], None),
                     ("E6 (mémoire de requirements.md)", DEFAULTS["finance"], DEFAULTS["database"], mem_e6)]
        e7_fin = os.path.join(JEU, f"Xbh Contracts Financial performance - {mois} (FICTIF E7).xlsx")
        e7_db = os.path.join(JEU, "Database Hcare XBH in service contracts (FICTIF E7).xlsx")
        if os.path.exists(e7_fin) and os.path.exists(e7_db):
            scenarios.append(("E7 (multiservices, mémoire vide)", e7_fin, e7_db, None))
        out_dir = runs_dir or os.path.join(tmp, "runs")
        for label, fin, db, mem in scenarios:
            try:
                results.append(run(fin, db, DEFAULTS["pseudo"], mem, mois, out_dir, label))
            except Stop as e:
                results.append({"scenario": label, "arret": str(e)})
        # Étape 8 de bout en bout : E5, décisions de Baptiste, puis relance avec la mémoire obtenue.
        label = "E5 → décisions (FB.TEST060 = Oui, FH.TEST011 = Non) → relance"
        try:
            mem_chain = os.path.join(tmp, "memoire-chaine.csv")
            first = run(DEFAULTS["finance"], DEFAULTS["database"], DEFAULTS["pseudo"], None, mois, out_dir,
                        label + " [1. import]")
            after = apply_decisions(first["fichier"], mem_chain, {"FB.TEST060": "Oui", "FH.TEST011": "Non"})
            with open(mem_chain, newline="", encoding="utf-8") as f:
                after["memoire_contenu"] = [r for r in csv.DictReader(f)]
            second = run(DEFAULTS["finance"], DEFAULTS["database"], DEFAULTS["pseudo"], mem_chain, mois,
                         out_dir, label + " [3. relance]")
            results.append({"scenario": label, "apres_decisions": after,
                            "relance": {"fichier": second["fichier"], "counts": second["counts"],
                                        "a_confirmer": second["a_confirmer"], "controle": second["controle"]}})
        except Stop as e:
            results.append({"scenario": label, "arret": str(e)})
    return results


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Simulation de l'import mensuel finance (données fictives).")
    ap.add_argument("usage", choices=["simulation", "decisions", "tests"])
    ap.add_argument("--fichier", help="fichier du mois produit par l'import (usage decisions)")
    ap.add_argument("--decision", action="append", default=[],
                    help="décision simulée, ex. FB.TEST060=Oui (répétable ; sinon la colonne du fichier fait foi)")
    ap.add_argument("--garder", action="store_true", help="tests : écrire les fichiers dans runs/")
    ap.add_argument("--mois", default="Octobre 2026")
    ap.add_argument("--finance", default=DEFAULTS["finance"])
    ap.add_argument("--database", default=DEFAULTS["database"])
    ap.add_argument("--pseudo", default=DEFAULTS["pseudo"])
    ap.add_argument("--memoire", default=DEFAULTS["memoire"])
    ap.add_argument("--runs", default=RUNS)
    a = ap.parse_args()
    try:
        if a.usage == "tests":
            out = run_tests(a.mois, a.runs if a.garder else None)
        elif a.usage == "decisions":
            if not a.fichier:
                raise Stop("Indiquer le fichier du mois avec --fichier.")
            forced = None
            if a.decision:
                forced = {}
                for d in a.decision:
                    k, _, v = d.partition("=")
                    if v.strip().lower() not in ("oui", "non"):
                        raise Stop(f"Décision invalide « {d} » : attendu CLÉ=Oui ou CLÉ=Non.")
                    forced[norm(k)] = v.strip().capitalize()
            out = apply_decisions(a.fichier, a.memoire, forced)
        else:
            out = run(a.finance, a.database, a.pseudo, a.memoire, a.mois, a.runs, "simulation")
        print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    except Stop as e:
        print(json.dumps({"arret": str(e)}, ensure_ascii=False))
        sys.exit(2)
